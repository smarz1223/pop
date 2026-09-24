"""
Pop Never Played Fantasy Football - data pipeline
Runs daily on GitHub Actions. Downloads the two published Google Sheets
workbooks, scores every player, reconciles to Weekly Scores, builds
standings (H2H only, no median game), analytics, and league history.

Outputs:
  data/pop_data.json      everything the site needs
  data/recon_report.md    human-readable reconciliation report

Local test: set env vars STATS_FILE and HISTORY_FILE to .xlsx paths.
"""
import io, json, math, os, re, sys, collections, datetime
import openpyxl

# ----------------------------------------------------------------- CONFIG
STATS_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vSt0W9qOb-u-BMpCYn1Ujj54kC2-kY35A-SEpxXfbF-C1-Diaiu8nkaoM-ICnDOD6vAnHYeVgZ_ZMs2/pub?output=xlsx")
HISTORY_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vSeKmWG0pwmPah071Ebu-UDFqH5UWjmcfjngYEKQl2V-ZhGdIhYdTv7btRn_N4dzHIfNZBlwTl2cnzp/pub?output=xlsx")

SEASON = 2026
REG_SEASON_WEEKS = 14          # regular season weeks 1-14
PLAYOFF_WEEKS = [15, 16, 17]
PLAYOFF_TEAMS = 6
BYE_TEAMS = 2
OUT_DIR = "data"

# History names -> 2026 names (2026 names display everywhere)
HISTORY_NAME_MAP = {"Vin": "VINNY", "Jeeps": "LITTLE JAMES", "Little David": "LITTLE BUBS",
                    "Marz": "STEVE"}

# 2026 owner renames, applied to Weekly Scores and the game log tables.
# Keeps the site right whether or not the workbook itself has been updated.
OWNER_RENAME = {"MARZ": "STEVE"}
LEAGUE_SIZE = 12

# League scoring (from SCORING MODIFIERS tab)
SCORING = {
    # offense
    "Pass Yds": 1 / 25, "Pass TD": 4, "Pass INT": -1,
    "Rush Yds": 1 / 10, "Rush TD": 6,
    "Receptions": 0.5, "Rec Yds": 1 / 10, "Rec TD": 6,
    "Return TD": 6, "2-Pt Conv": 2, "Fumbles Lost": -2,
    # kicker
    "FG 0-19": 3, "FG 20-29": 3, "FG 30-39": 3, "FG 40-49": 4, "FG 50+": 5,
    "PAT": 1,
    # defense / special teams
    "Sacks": 1, "Def INT": 2, "Fum Rec": 2, "Def TD": 6, "Safety": 2,
    "Blocked Kick": 2, "Def Return TD": 6,
}
DST_PA_TIERS = [10, 7, 4, 1, 0, -1, -4]   # points-allowed tier values

# Column map: (group row label, header row label) -> stat key
POS_COLS = {("Misc", "GP*"): "GP", ("Misc", "2PT"): "2-Pt Conv",
            ("Passing", "Yds"): "Pass Yds", ("Passing", "TD"): "Pass TD", ("Passing", "Int"): "Pass INT",
            ("Rushing", "Yds"): "Rush Yds", ("Rushing", "TD"): "Rush TD",
            ("Receiving", "Rec"): "Receptions", ("Receiving", "Yds"): "Rec Yds", ("Receiving", "TD"): "Rec TD",
            ("Ret", "TD"): "Return TD", ("Fum", "Lost"): "Fumbles Lost"}
K_COLS = {("Misc", "GP*"): "GP",
          ("Field Goals Made", "0-19"): "FG 0-19", ("Field Goals Made", "20-29"): "FG 20-29",
          ("Field Goals Made", "30-39"): "FG 30-39", ("Field Goals Made", "40-49"): "FG 40-49",
          ("Field Goals Made", "50+"): "FG 50+",
          ("PAT", "Made"): "PAT"}
DST_COLS = {("Misc", "GP*"): "GP", ("Misc", "Blk Kick"): "Blocked Kick",
            ("Tackles", "Sack"): "Sacks", ("Tackles", "Safe"): "Safety",
            ("Turnovers", "Int"): "Def INT", ("Turnovers", "Fum Rec"): "Fum Rec",
            ("TD", "TD"): "Def TD", ("Ret", "TD"): "Def Return TD"}

# Stat category groupings for the Stat Categories page
CATEGORY_GROUPS = {
    "Passing": ["Pass Yds", "Pass TD", "Pass INT"],
    "Rushing": ["Rush Yds", "Rush TD"],
    "Receiving": ["Receptions", "Rec Yds", "Rec TD"],
    "Misc Offense": ["Return TD", "2-Pt Conv", "Fumbles Lost"],
    "Kicking": ["FG 0-19", "FG 20-29", "FG 30-39", "FG 40-49", "FG 50+", "PAT"],
    "Defense": ["Sacks", "Def INT", "Fum Rec", "Def TD", "Safety", "Blocked Kick",
                "Def Return TD", "DST Pts Allowed"],
}
POSITIONS = ["QB", "RB", "WR", "TE", "K", "DEF"]

# Yahoo name-cell cleanup
NOTE_FLAGS = ["No new player Notes", "No new player Note", "New Player Notes", "New Player Note",
              "Player Notes", "Player Note", "Video Forecast"]
STATUS_TAGS = sorted(["IR-R", "IR", "PUP-P", "PUP-R", "NFI-R", "NFI-A", "SUSP", "COVID-19",
                      "DTD", "NA", "O", "Q", "D"], key=len, reverse=True)
TEAM_POS_RE = re.compile(r"([A-Z][A-Za-z]{1,2}) - (QB|WR|RB|TE|K|DEF)\s*$")


# ----------------------------------------------------------------- LOADING
def load_workbook(env_var, url):
    path = os.environ.get(env_var)
    if path:
        return openpyxl.load_workbook(path, data_only=True)
    import requests
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def r2(x):
    return round(x + 0.0, 2)


def parse_name(raw):
    first = str(raw).split("\n")[0].strip()
    m = TEAM_POS_RE.search(first)
    nfl, pos = (m.group(1).upper(), m.group(2)) if m else (None, None)
    name = first[:m.start()] if m else first
    for f in NOTE_FLAGS:
        name = name.replace(f, "")
    name = name.strip()
    for tag in STATUS_TAGS:  # status glued to the name, e.g. "A.J. BrownIR"
        if name.endswith(tag) and len(name) > len(tag):
            prev = name[-len(tag) - 1]
            if prev.islower() or prev in ".'":
                name = name[:-len(tag)].strip()
                break
    return name, nfl, pos


def read_log_table(ws, colmap, kind):
    """Returns (player_rows, yahoo_totals_by_owner)."""
    rows = list(ws.iter_rows(values_only=True))
    groups, cur = [], None
    for v in rows[0]:
        if v not in (None, ""):
            cur = str(v).strip()
        groups.append(cur)
    headers = [str(v).strip() if v is not None else None for v in rows[1]]
    idx = {}
    for i, (g, h) in enumerate(zip(groups, headers)):
        if (g, h) in colmap:
            idx[colmap[(g, h)]] = i
    missing = set(colmap.values()) - set(idx)
    if missing:
        raise RuntimeError(f"{ws.title}: columns not found {missing}")
    players, totals = [], {}
    for r in rows[2:]:
        owner, name = r[0], r[2]
        if not owner or not name:
            continue
        owner = OWNER_RENAME.get(str(owner).strip(), str(owner).strip())
        stats = {k: num(r[i]) for k, i in idx.items()}
        if str(name).strip() == "Totals":
            totals[owner] = stats
            continue
        pname, nfl, pos = parse_name(name)
        players.append({"owner": owner, "player": pname, "nfl": nfl,
                        "pos": pos, "kind": kind, "stats": stats})
    return players, totals


def score(stats):
    pts = {k: stats.get(k, 0) * v for k, v in SCORING.items() if k in stats}
    return pts


# ----------------------------------------------------------------- WEEKLY / STANDINGS
def load_weekly(ws):
    """Finds the header row by name, so the TEAM NAME column (or any
    column reshuffle) can't shift the scores."""
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:10])
              if any(str(v).strip() == "Points For" for v in r if v is not None))
    col = {str(v).strip(): i for i, v in enumerate(rows[hi]) if v is not None}
    iw, it, ipf, ipa = col["Week"], col["Team"], col["Points For"], col["Points Against"]
    ren = lambda t: OWNER_RENAME.get(t, t)
    weeks = collections.defaultdict(dict)
    for r in rows[hi + 1:]:
        if len(r) <= max(iw, it, ipf, ipa):
            continue
        wk, team, pf, pa = r[iw], r[it], r[ipf], r[ipa]
        if wk is None or not team or pf in (None, "") or pa in (None, ""):
            continue
        weeks[int(wk)][ren(str(team).strip())] = {"pf": float(pf), "pa": float(pa)}
    return dict(sorted(weeks.items()))


def pair_opponents(weeks, flags):
    for wk, teams in weeks.items():
        for t, d in teams.items():
            cands = [o for o, e in teams.items() if o != t
                     and abs(e["pf"] - d["pa"]) < 0.005 and abs(e["pa"] - d["pf"]) < 0.005]
            if len(cands) == 1:
                d["opp"] = cands[0]
            else:
                d["opp"] = None
                flags.append(f"Week {wk}: could not uniquely match opponent for {t} ({len(cands)} candidates)")


def result(a, b):
    return "W" if a > b + 1e-9 else ("L" if a < b - 1e-9 else "T")


def norm_cdf(z):
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def build_standings(weeks, owners):
    rec = {o: collections.Counter() for o in owners}
    weekly_rows = []
    perf = collections.defaultdict(list)
    for wk, teams in weeks.items():
        scores = sorted((d["pf"] for d in teams.values()), reverse=True)
        n = len(scores)
        mean = sum(scores) / n
        sd = math.sqrt(sum((s - mean) ** 2 for s in scores) / n) or 1
        reg = wk <= REG_SEASON_WEEKS
        for t, d in teams.items():
            h2h = result(d["pf"], d["pa"])
            rank = 1 + sum(1 for s in scores if s > d["pf"] + 1e-9)
            ap = collections.Counter(result(d["pf"], e["pf"]) for o, e in teams.items() if o != t)
            weekly_rows.append({"week": wk, "team": t, "opp": d.get("opp"), "pf": r2(d["pf"]),
                                "pa": r2(d["pa"]), "h2h": h2h, "avg": r2(mean), "rank": rank,
                                "allplay": f"{ap['W']}-{ap['L']}-{ap['T']}",
                                "playoff": not reg})
            if not reg:
                continue
            c = rec[t]
            c["H2H_" + h2h] += 1
            c["AP_W"] += ap["W"]; c["AP_L"] += ap["L"]; c["AP_T"] += ap["T"]
            c["PF"] += d["pf"]; c["PA"] += d["pa"]; c["G"] += 1
            perf[t].append(norm_cdf((d["pf"] - mean) / sd) * 100)
    table = []
    for t in owners:
        c = rec[t]
        W, L, T = c["H2H_W"], c["H2H_L"], c["H2H_T"]
        games = W + L + T
        h2h_g = c["H2H_W"] + c["H2H_L"] + c["H2H_T"]
        ap_g = c["AP_W"] + c["AP_L"] + c["AP_T"]
        h2h_pct = (c["H2H_W"] + 0.5 * c["H2H_T"]) / h2h_g if h2h_g else 0
        ap_pct = (c["AP_W"] + 0.5 * c["AP_T"]) / ap_g if ap_g else 0
        table.append({
            "team": t, "W": W, "L": L, "T": T,
            "win_pct": round((W + 0.5 * T) / games, 4) if games else 0,
            "h2h": f"{c['H2H_W']}-{c['H2H_L']}-{c['H2H_T']}",
            "allplay": f"{c['AP_W']}-{c['AP_L']}-{c['AP_T']}",
            "allplay_pct": round(ap_pct, 4),
            "exp_h2h_wins": round(ap_pct * h2h_g, 2),          # all-play expected H2H wins
            "luck": round((h2h_pct - ap_pct + 1) * 50, 1),       # 0-100, 50 = neutral
            "perf_rating": round(sum(perf[t]) / len(perf[t]), 1) if perf[t] else None,
            "PF": r2(c["PF"]), "PA": r2(c["PA"]),
            "PPG": r2(c["PF"] / c["G"]) if c["G"] else 0, "G": c["G"],
        })
    table.sort(key=lambda x: (-x["win_pct"], -x["PF"]))
    for i, row in enumerate(table, 1):
        row["seed"] = i
        row["status"] = "bye" if i <= BYE_TEAMS else ("playoff" if i <= PLAYOFF_TEAMS else "out")
    return table, weekly_rows


# ----------------------------------------------------------------- RECONCILIATION
def achievable_pa_sums(n):
    sums = {0}
    for _ in range(int(n)):
        sums = {s + t for s in sums for t in DST_PA_TIERS}
    return sums


def reconcile(players, yahoo_totals, weeks, owners, flags):
    report = []
    # 1. parsed rows vs Yahoo "Totals" rows
    for kind, tot in yahoo_totals.items():
        for o, ytot in tot.items():
            mine = collections.Counter()
            for p in players:
                if p["owner"] == o and p["kind"] == kind:
                    mine.update(p["stats"])
            for k, v in ytot.items():
                if abs(mine[k] - v) > 0.001:
                    flags.append(f"{o} {kind} table: parsed {k}={mine[k]} vs Yahoo Totals row {v}")
    # 2. freshness: game log vs weeks entered
    weeks_entered = len(weeks)
    max_gp = max((p["stats"].get("GP", 0) for p in players), default=0)
    state = "OK"
    if max_gp > weeks_entered:
        state = "LIVE"
        flags.append(f"Game log shows {int(max_gp)} games but Weekly Scores has {weeks_entered} weeks. "
                     "Week in progress; reconciliation will settle once scores are entered.")
    elif max_gp < weeks_entered:
        state = "STALE"
        flags.append(f"Game log shows only {int(max_gp)} games but Weekly Scores has {weeks_entered} weeks. "
                     "IMPORTHTML may not have refreshed.")
    # 3. per-owner: residual must be a valid DST points-allowed total
    dst_pa = {}
    for o in owners:
        wk_total = sum(w[o]["pf"] for w in weeks.values() if o in w)
        calc = sum(p["points"] for p in players if p["owner"] == o)
        dst_gp = sum(p["stats"].get("GP", 0) for p in players if p["owner"] == o and p["kind"] == "DST")
        resid = wk_total - calc
        ok = abs(resid - round(resid)) < 0.011 and round(resid) in achievable_pa_sums(dst_gp)
        if state == "OK" and not ok:
            flags.append(f"{o}: residual {resid:.2f} is not a valid DST points-allowed total "
                         f"for {int(dst_gp)} DST games")
        dst_pa[o] = round(resid)
        report.append({"team": o, "weekly_total": r2(wk_total), "calculated": r2(calc),
                       "dst_pts_allowed": round(resid), "dst_games": int(dst_gp),
                       "check": "PASS" if ok else ("PENDING" if state != "OK" else "FAIL")})
    overall = "FAIL" if any(r["check"] == "FAIL" for r in report) or \
        any("Totals row" in f or "opponent" in f for f in flags) else state
    if overall == "OK":
        overall = "PASS"
    return report, dst_pa, overall, weeks_entered


# ----------------------------------------------------------------- ANALYTICS
def build_analytics(players, dst_pa, owners):
    raw = {o: collections.Counter() for o in owners}
    pts = {o: collections.Counter() for o in owners}
    by_pos = {o: collections.Counter() for o in owners}
    for p in players:
        o = p["owner"]
        raw[o].update({k: v for k, v in p["stats"].items() if k != "GP"})
        pts[o].update(p["pts_by_cat"])
        by_pos[o][p["pos"]] += p["points"]
    for o in owners:
        pts[o]["DST Pts Allowed"] += dst_pa.get(o, 0)
        by_pos[o]["DEF"] += dst_pa.get(o, 0)
    cats = {}
    for group, keys in CATEGORY_GROUPS.items():
        cats[group] = {o: {"raw": {k: raw[o].get(k, 0) for k in keys if k != "DST Pts Allowed"},
                           "pts": {k: r2(pts[o].get(k, 0)) for k in keys},
                           "total": r2(sum(pts[o].get(k, 0) for k in keys))} for o in owners}
    positions = {o: {pos: r2(by_pos[o].get(pos, 0)) for pos in POSITIONS} for o in owners}
    return cats, positions


# ----------------------------------------------------------------- HISTORY
CURRENT_OWNERS = set()


def hname(n):
    """History name -> display name. Current managers use their 2026 name."""
    n = str(n).strip()
    n = HISTORY_NAME_MAP.get(n, n)
    return n.upper() if n.upper() in CURRENT_OWNERS else n


def build_history(wb, flags):
    ws = wb["Data"]
    hdr = [str(v).strip() if v else None for v in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(hdr) if h}
    seasons = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Years"]] in (None, ""):
            continue
        g = lambda k: num(r[col[k]])
        seasons.append({"year": int(r[col["Years"]]), "manager": hname(r[col["Manager"]]),
                        "active": str(r[col["Status"]]).strip() == "Active",
                        "first": g("1st"), "second": g("2nd"), "third": g("3rd"),
                        "W": g("Wins"), "L": g("Losses"), "T": g("Ties"),
                        "games": g("Games"), "playoffs": g("Playoffs")})
    # all-time
    agg = collections.defaultdict(collections.Counter)
    active = {}
    for s in seasons:
        a = agg[s["manager"]]
        a["years"] += 1
        for k in ("first", "second", "third", "W", "L", "T", "games", "playoffs"):
            a[k] += s[k]
        active[s["manager"]] = active.get(s["manager"], False) or s["active"]
    alltime = []
    for m, a in agg.items():
        alltime.append({"manager": m, "active": active[m], "years": a["years"], "games": a["games"],
                        "first": a["first"], "second": a["second"], "third": a["third"],
                        "top2": a["first"] + a["second"], "top3": a["first"] + a["second"] + a["third"],
                        "W": a["W"], "L": a["L"], "T": a["T"],
                        "win_pct": round((a["W"] + 0.5 * a["T"]) / a["games"], 4) if a["games"] else 0,
                        "playoffs": a["playoffs"],
                        "playoff_pct": round(a["playoffs"] / a["years"], 4)})
    alltime.sort(key=lambda x: (not x["active"], -x["win_pct"]))
    # active rankings (1 = best; losses: fewest = best)
    act = [x for x in alltime if x["active"]]
    ranks = {x["manager"]: {} for x in act}
    for k in ("years", "games", "first", "second", "third", "top2", "top3", "W", "L",
              "win_pct", "playoffs", "playoff_pct"):
        for x in act:
            better = sum(1 for y in act if (y[k] < x[k] if k == "L" else y[k] > x[k]))
            ranks[x["manager"]][k] = better + 1
    # yearly grids
    years = sorted({s["year"] for s in seasons})
    wins_grid = {y: {} for y in years}
    playoff_grid = {y: {} for y in years}
    for s in seasons:
        wins_grid[s["year"]][s["manager"]] = s["W"]
        playoff_grid[s["year"]][s["manager"]] = int(s["playoffs"])
    games_by_year = {y: max(s["games"] for s in seasons if s["year"] == y) for y in years}
    # championships: champion / runner-up from Data; draft slot from the
    # Summary tab's CHAMPION'S SEASON block (rows are in year order)
    draft_by_year = {}
    if "Summary" in wb.sheetnames:
        sm = wb["Summary"]
        hdr_cell = next((c for row in sm.iter_rows(max_row=15) for c in row
                         if str(c.value or "").strip() == "Draft #"), None)
        if hdr_cell:
            vals = []
            for rr in range(hdr_cell.row + 1, hdr_cell.row + 1 + len(years)):
                v = sm.cell(rr, hdr_cell.column).value
                w = sm.cell(rr, hdr_cell.column - 3).value
                l = sm.cell(rr, hdr_cell.column - 2).value
                vals.append((v, w, l))
            for y, (v, w, l) in zip(years, vals):
                if v in (None, ""):
                    continue
                c = next((s for s in seasons if s["year"] == y and s["first"] == 1), None)
                if c and (abs(num(w) - c["W"]) > 0.01 or abs(num(l) - c["L"]) > 0.01):
                    flags.append(f"History {y}: Summary champion record ({num(w):g}-{num(l):g}) "
                                 f"differs from Data ({c['W']:g}-{c['L']:g}); draft slot skipped")
                    continue
                draft_by_year[y] = int(num(v))
    champs = []
    for y in years:
        c = next((s for s in seasons if s["year"] == y and s["first"] == 1), None)
        ru = next((s for s in seasons if s["year"] == y and s["second"] == 1), None)
        champs.append({"year": y, "champion": c["manager"] if c else None,
                       "runner_up": ru["manager"] if ru else None,
                       "W": c["W"] if c else 0, "L": c["L"] if c else 0, "T": c["T"] if c else 0,
                       "draft": draft_by_year.get(y),
                       "size": sum(1 for s in seasons if s["year"] == y)})
    slot_titles = collections.Counter(c["draft"] for c in champs
                                      if c["size"] == LEAGUE_SIZE and c["draft"])
    # evolution: seats keep their position year to year; a new manager takes
    # the seat of the one who left
    members = {y: [s["manager"] for s in seasons if s["year"] == y] for y in years}
    evolution, seats_prev, seen = [], [], set()
    for i, y in enumerate(years):
        cur = members[y]
        seats = [m if m in cur else None for m in seats_prev]
        for m in cur:
            if m not in seats:
                if None in seats:
                    seats[seats.index(None)] = m
                else:
                    seats.append(m)
        nxt = set(members[years[i + 1]]) if i + 1 < len(years) else None
        prev = set(members[years[i - 1]]) if i else set()
        row = []
        for m in seats:
            if m is None:
                row.append(None)
                continue
            if nxt is not None and m not in nxt:
                t = "Departing"
            elif i == 0:
                t = "Original"
            elif m not in seen:
                t = "New"
            elif m not in prev:
                t = "Returning"
            else:
                t = "Continuing"
            row.append({"manager": m, "type": t})
        seen.update(cur)
        seats_prev = seats
        evolution.append({"year": y, "seats": row})
    # awards (recomputed)
    def best(key, rows, rev=True, fmt=lambda v: v):
        if not rows:
            return None
        top = max(r[key] for r in rows) if rev else min(r[key] for r in rows)
        who = [r for r in rows if abs(r[key] - top) < 1e-9]
        return top, who
    awards = []
    def add(title, key, rows, rev=True, label=lambda r: r["manager"], extra=lambda r: "All"):
        res = best(key, rows, rev)
        if res:
            v, who = res
            awards.append({"award": title, "stat": v, "managers": sorted({label(w) for w in who}),
                           "years": sorted({str(extra(w)) for w in who})})
    add("Most Titles", "first", alltime)
    add("Most Wins", "W", alltime)
    add("Most Losses", "L", alltime)
    add("Best Win % (2+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 2])
    add("Worst Win % (2+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 2], rev=False)
    add("Most Playoff Appearances", "playoffs", alltime)
    add("Most Top 2 Finishes", "top2", alltime)
    add("Most Top 3 Finishes", "top3", alltime)
    sp = [dict(s, pct=(s["W"] + 0.5 * s["T"]) / s["games"]) for s in seasons if s["games"]]
    add("Best Single Season (Win %)", "pct", sp, extra=lambda r: r["year"])
    add("Worst Single Season (Win %)", "pct", sp, rev=False, extra=lambda r: r["year"])
    return {"alltime": alltime, "active_ranks": ranks, "years": years,
            "games_by_year": games_by_year,
            "wins_grid": wins_grid, "playoff_grid": playoff_grid,
            "championships": champs, "titles_by_draft_slot": dict(sorted(slot_titles.items())),
            "evolution": evolution, "awards": awards, "seasons": seasons}


# ----------------------------------------------------------------- MAIN
def main():
    flags = []
    swb = load_workbook("STATS_FILE", STATS_URL)
    hwb = load_workbook("HISTORY_FILE", HISTORY_URL)

    weeks = load_weekly(swb["Weekly Scores"])
    owners = list(dict.fromkeys(t for w in weeks.values() for t in w))
    pair_opponents(weeks, flags)

    players, ytot = [], {}
    for sheet, cmap, kind in (("POSITION_TABLE", POS_COLS, "OFF"),
                              ("KICKER_TABLE", K_COLS, "K"), ("DST_TABLE", DST_COLS, "DST")):
        p, t = read_log_table(swb[sheet], cmap, kind)
        players += p
        ytot[kind] = t
    for p in players:
        p["pts_by_cat"] = score(p["stats"])
        p["points"] = sum(p["pts_by_cat"].values())
        if not p["pos"]:
            flags.append(f"Could not read position for '{p['player']}' ({p['owner']})")
    unknown = {p["owner"] for p in players} - set(owners)
    for u in sorted(unknown):
        flags.append(f"Game log owner '{u}' not found in Weekly Scores")

    recon, dst_pa, status, weeks_entered = reconcile(players, ytot, weeks, owners, flags)
    standings, weekly_rows = build_standings(weeks, owners)
    cats, positions = build_analytics(players, dst_pa, owners)
    CURRENT_OWNERS.update(owners)
    history = build_history(hwb, flags)

    player_out = sorted(({"owner": p["owner"], "player": p["player"], "nfl": p["nfl"], "pos": p["pos"],
                          "gp": int(p["stats"].get("GP", 0)), "points": r2(p["points"]),
                          "stats": {k: v for k, v in p["stats"].items() if k != "GP"},
                          "pts_by_cat": {k: r2(v) for k, v in p["pts_by_cat"].items() if v}}
                         for p in players), key=lambda x: -x["points"])

    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    out = {"season": SEASON, "updated": now, "weeks_entered": weeks_entered,
           "reg_season_weeks": REG_SEASON_WEEKS, "playoff_weeks": PLAYOFF_WEEKS,
           "owners": owners, "recon_status": status, "recon": recon, "flags": flags,
           "standings": standings, "weekly": weekly_rows, "categories": cats,
           "positions": positions, "players": player_out, "history": history,
           "scoring": SCORING}
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "pop_data.json"), "w") as f:
        json.dump(out, f, indent=1, default=float)

    lines = [f"# Pop Reconciliation: {status}", f"Updated {now} | Weeks entered: {weeks_entered}", "",
             "| Team | Weekly Total | Calculated | DST Pts Allowed | DST Games | Check |",
             "|---|---|---|---|---|---|"]
    for r in recon:
        lines.append(f"| {r['team']} | {r['weekly_total']:.2f} | {r['calculated']:.2f} | "
                     f"{r['dst_pts_allowed']} | {r['dst_games']} | {r['check']} |")
    lines += ["", "## Flags"] + ([f"- {f}" for f in flags] or ["- None"])
    with open(os.path.join(OUT_DIR, "recon_report.md"), "w") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines))
    # never fail the run on data flags; the site shows the status badge instead


if __name__ == "__main__":
    main()
